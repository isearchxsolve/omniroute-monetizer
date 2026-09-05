from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\omniroute-monetizer\state\Laptop_Phone_Network_Health_Report_2026_09_05.docx")

NAVY = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=105, bottom=90, end=105):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def write_cell(cell, text, bold=False, color=BLACK, size=9.3, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_table_widths(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        write_cell(cell, text, bold=True, color=RGBColor(255, 255, 255), size=9.1,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_idx, row_data in enumerate(rows):
        row = table.add_row()
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            if row_idx % 2 == 1:
                set_cell_shading(cell, LIGHT_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            write_cell(cell, str(text), size=9.1, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.bold = True
        lead.font.color.rgb = BLACK
        lead.font.name = "Aptos"
        lead._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        lead._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run = p.add_run(text)
    run.font.color.rgb = BLACK
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        run = p.add_run(item)
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.color.rgb = BLACK
    run.font.name = "Aptos Display" if level == 1 else "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), run.font.name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), run.font.name)
    return p


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = BLACK
    for name, size in (("Title", 22), ("Subtitle", 11), ("Heading 1", 14), ("Heading 2", 11)):
        style = styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = name != "Subtitle"
        if name == "Title":
            # Word's stock Title style carries a blue bottom border. This
            # report deliberately uses a plain black title with no rule.
            ppr = style.element.get_or_add_pPr()
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)


def build_report():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Laptop Phone and Network Health Report")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Diagnostic Findings and Practical Care Plan")
    subtitle.paragraph_format.space_after = Pt(13)

    add_body(doc, "This report records the completed diagnostics and the practical decisions for the laptop, audio, network, and phone. The main conclusion is that the laptop remains usable for work when plugged into stable power and managed carefully. The most important risks are the worn laptop battery, limited free system-drive space, the old but presently stable Wi-Fi adapter, weak built-in speakers, and the phone's prior heat and frame-opening history.")

    add_heading(doc, "At a Glance")
    add_table(doc, ["Area", "Current Position", "What To Do"], [
        ("Laptop", "Usable on AC power with no current disk, NTFS, SATA, WHEA, or bugcheck evidence of an imminent failure", "Keep work backed up and avoid treating it as the only copy of important work"),
        ("Laptop battery", "About 11.6 percent of original design capacity remains; observed estimate about 39 minutes, much less during heavy work", "Use AC power for work and treat battery time as an emergency bridge only"),
        ("System drive", "C drive improved from 11.5 GB free to 25.33 GB free after safe cleanup and hibernation removal", "Keep at least 25 GB free and avoid large unnecessary downloads on C drive"),
        ("External microphone", "External Realtek jack microphone is now the Windows default at 80 percent and plus 10.04 dB", "Use the external microphone for calls; do not return it to 100 percent boost"),
        ("Built in speakers", "Windows levels are already at maximum; headphones sound normal", "Use headphones for clear audio; internal speakers are likely physically weak"),
        ("Airtel Wi Fi", "5 GHz link tested with 0 percent loss and 33 ms average internet latency", "Use as the primary connection for normal work"),
        ("Phone tethering", "USB tethering works but Jio mobile latency was 54 to 223 ms with 109 ms average", "Use only as a short attended backup if Wi Fi fails"),
        ("Phone safety", "No active warning signs reported recently, but there is a prior heat and frame-opening history", "Avoid continuous charging and heat; get a physical inspection if the gap changes or symptoms return"),
    ], [1.25, 2.78, 2.78])

    add_heading(doc, "Laptop Reliability and Storage")
    add_body(doc, "The laptop is a Lenovo ThinkPad T430 with an Intel Core i5 3320M processor, 16 GB of memory, and a 480 GB Intel SSD. The machine is old, but the diagnostic did not find current Windows evidence of an imminent CPU, RAM, SSD, disk-controller, or filesystem failure. That is useful evidence, not a warranty: old fans, charger connectors, hinges, display cables, and the motherboard can still fail without warning.")
    add_table(doc, ["Finding", "Measured Result", "Meaning"], [
        ("Battery health", "6,710 mWh full charge capacity versus 57,720 mWh design capacity", "The original laptop battery is heavily worn and should not be relied upon for a work session"),
        ("Unexpected shutdown history", "Five Kernel Power 41 events in 30 days; bugcheck count 0", "The system has experienced abrupt power loss or forced shutdowns, not proven blue-screen failures"),
        ("Hibernate reliability", "87 earlier S4 resume firmware warnings; hibernation disabled with authorized Windows command", "Fast Startup and hibernation are off; normal S3 sleep remains available"),
        ("C drive space", "25.33 GB free after safe cache cleanup and removing a 6.28 GB hibernation file", "The emergency low-space condition was reduced, but free space still needs regular attention"),
        ("SSD health", "Windows reported Healthy and OK; detailed SMART wear and temperature were unavailable", "No current operating-system disk warning; detailed remaining SSD life could not be certified"),
    ], [1.35, 2.55, 2.91])
    add_heading(doc, "Laptop Care Plan", 2)
    add_bullets(doc, [
        "Use the laptop plugged into a stable charger for long work. Do not plan important work around the worn internal battery.",
        "Keep regular saves and remote backups. The source repository was confirmed synchronized with its remote copy during the work.",
        "Keep the cooling pad running. Do not block its vents or place the laptop on bedding or a soft surface.",
        "Keep the C drive above the present 25 GB free-space margin. Store large personal files on D drive or an external destination when possible.",
        "Avoid several heavy browser, AI, video, and build tasks at once. The old dual-core processor is the performance ceiling.",
    ])

    add_heading(doc, "Microphone and Speaker Results")
    add_body(doc, "The original microphone problem came from maximum Realtek capture gain and the laptop cooler adding steady fan and vibration noise. The final working input is the external microphone in the round 3.5 mm port, not the laptop microphone array.")
    add_table(doc, ["Audio Item", "Verified Result", "Current Recommendation"], [
        ("External microphone", "Detected at the Realtek microphone jack; set as Console, Multimedia, and Communications default", "Use Microphone Realtek High Definition Audio inside apps that have their own input selector"),
        ("External microphone gain", "Reduced from 100 percent and plus 22.50 dB to 80 percent and plus 10.04 dB", "Keep 80 percent; speech was captured without clipping and with about 18.6 dB peak headroom"),
        ("Internal microphone array", "Cooler-on background noise fell about 18 dB after lowering maximum boost", "Retained only as a fallback; it is no longer the Windows default"),
        ("Built in speakers", "Correct endpoint, 100 percent master, 100 percent left and right balance, no mute, no device error", "Do not use volume booster software; use wired headphones for clear sound"),
        ("Headphone comparison", "Operator hears normal sound through headphones", "Confirms the Realtek audio path works and internal speakers are the weak component"),
    ], [1.42, 2.9, 2.49])
    add_body(doc, "The built-in speaker problem is most consistent with old, quiet physical speaker drivers, obstructed grilles, or aging speaker wiring. Windows volume and driver routing are already correct. A future technician should inspect or replace the internal speaker assembly, not replace the sound card without evidence.")

    add_heading(doc, "Network Results")
    add_body(doc, "The normal Wi-Fi connection is working well now. The ThinkPad uses an Intel Centrino Advanced N 6205 adapter, a legacy 802.11n part with a signed 2015 driver. The adapter has occasionally logged driver-originated disconnects in history, but it performed cleanly during the live tests and no newer current Intel driver download was identified.")
    add_table(doc, ["Connection", "Measured Result", "Decision"], [
        ("Primary Airtel 5 GHz Wi-Fi", "80 percent signal, 240 Mbps negotiated link, 0 percent loss to router and internet, 33 ms average to 1.1.1.1", "Keep as the primary connection"),
        ("Laptop Wi-Fi adapter", "Transmit power Highest; 802.11n enabled; no packet errors or discards in test", "Do not reset or change driver while it is working"),
        ("Jio phone USB tethering", "USB RNDIS connection worked with 0 percent loss; internet latency 54 to 223 ms, 109 ms average", "Emergency backup only, short and attended"),
        ("Tethering route behavior", "Windows preferred the phone route while tethering was enabled", "Turn tethering off and unplug when the temporary backup is no longer needed"),
    ], [1.55, 3.1, 2.16])
    add_body(doc, "For an important long session, the safest order is Airtel Wi-Fi first, then an existing Ethernet cable if one is already available, then short attended USB tethering. The phone connection is functional but has higher latency variation and adds phone heat and mobile-data use.")

    add_heading(doc, "Phone Safety and Charging")
    add_body(doc, "The connected phone was detected correctly as an iQOO Z7 Pro 5G in ordinary Android USB mode. The cable, laptop port, and phone USB connection are working. Ordinary USB mode does not expose battery health, temperature, charging watts, or internal phone sensors, so those values were not claimed.")
    add_table(doc, ["Observation", "Assessment", "Safe Response"], [
        ("Prior heat and swelling concern", "No recent swelling, heat, rocking, odour, hissing, smoke, or growing gap reported after avoiding heat and continuous full charge", "This lowers concern for an active emergency but does not certify the battery healthy"),
        ("Left frame opening", "Single photo did not show obvious severe doming or screen lift; visible seam was inconclusive", "Do not press it shut, heat it, glue it, or disassemble the phone"),
        ("Charging method", "Laptop USB generally provides slower charging than a dedicated fast charger", "Short attended charging while phone is idle may reduce heat, but is not a battery repair"),
        ("High heat combination", "5G mobile data, Wi-Fi hotspot or tethering, charging, screen use, and warm surroundings can stack heat", "Do not run this combination continuously or unattended"),
    ], [1.55, 3.05, 2.21])
    add_heading(doc, "Phone Rules", 2)
    add_bullets(doc, [
        "Do not leave the phone charging overnight or all day merely because laptop USB charging is slow.",
        "When charging, keep the phone on a hard ventilated surface with hotspot and USB tethering turned off.",
        "Stop using and charging it if the frame gap grows, becomes central or uneven, the phone rocks, the display or back lifts, or heat, leakage, odour, hissing, smoke, or vapour appears.",
        "Do not puncture, compress, bend, freeze, glue, or open the phone. Do not put a suspected lithium-ion battery in household rubbish.",
        "When financially practical, request a battery and frame inspection at an iQOO or vivo authorized centre. Ask them to check the battery before simply gluing the frame closed.",
    ])

    add_heading(doc, "Daily Checklist")
    add_table(doc, ["Before Work", "During Work", "Stop and Check"], [
        ("Use AC power, Airtel Wi-Fi, 25 GB or more free on C drive, and the external microphone", "Keep cooler airflow open, use headphones for clear sound, save and sync important work regularly", "Laptop abruptly loses power, C drive drops near 15 GB free, Wi-Fi repeatedly drops, or any new disk/WHEA warning appears"),
        ("Check phone is flat, cool, and not tethering unless needed", "If charging the phone, keep it idle, uncovered, and attended", "Phone gets hot, frame gap grows, screen/back lifts, smells odd, hisses, leaks, smokes, or gives off vapour"),
    ], [2.18, 2.18, 2.45])

    add_heading(doc, "Evidence Limits and Sources")
    add_body(doc, "This report distinguishes measured findings from practical interpretation. Windows could not provide detailed SSD SMART wear, SSD temperature, CPU temperature, fan speed, phone battery health, phone battery temperature, or charging wattage. A healthy device status does not prove old physical speakers, batteries, cables, fans, or motherboard parts have their original life remaining.")
    add_body(doc, "Official sources consulted during the diagnostics: Intel legacy wireless support and Windows compatibility guidance for the Centrino Advanced N 6205; Lenovo ThinkPad wireless LAN support listing for the T430; and Fire and Rescue NSW lithium-ion battery safety guidance. These sources supported the decision not to install an unverified Wi-Fi driver and the phone-battery safety precautions.")
    add_body(doc, "Report date: 5 September 2026. This is a practical maintenance record, not a warranty, repair certificate, or substitute for a physical inspection where a safety warning is present.")

    doc.core_properties.title = "Laptop Phone and Network Health Report"
    doc.core_properties.subject = "Personal device health and safety diagnostic record"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = ""
    doc.save(OUT)


if __name__ == "__main__":
    build_report()
    print(OUT)
